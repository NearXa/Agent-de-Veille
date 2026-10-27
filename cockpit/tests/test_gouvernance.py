import json

from models import Lot, Recommandation
from services.gouvernance_service import DecisionRefusee, GouvernanceService
from services.matrices import classeur_xlsx, toutes
from tests.test_models import EXEMPLE_SPEC
from models import Exigence, Texte


class FakeNotion:
    def __init__(self):
        self.lots_mis_a_jour = []
        self.pages_creees = []

    def ecrire_decision_gouvernance(self, cadre, auteur, date_iso, commentaire, lots):
        self.lots_mis_a_jour = [l.id for l in lots]
        self.pages_creees.append((cadre, auteur, commentaire))
        return len(lots)


def test_decision_de_gouvernance_ecrite_dans_notion_et_decision_json(tmp_path):
    notion = FakeNotion()
    service = GouvernanceService(notion, "Neyl Kaci", output_dir=tmp_path)
    lots = [Lot(id="LA-0001", titre="Lot 1", notion_page_id="page-1"), Lot(id="LA-0002", titre="Lot 2")]
    reco = Recommandation(cadre_recommande="Hybride", texte_id="TX-0001", run_id="RUN-1")
    decision = service.enregistrer("TOGAF", "priorité à l'architecture", lots, reco)
    assert decision.cadre == "TOGAF" and decision.auteur == "Neyl Kaci" and decision.texte_id == "TX-0001"
    assert notion.lots_mis_a_jour == ["LA-0001", "LA-0002"]
    assert notion.pages_creees == [("TOGAF", "Neyl Kaci", "priorité à l'architecture")]
    saved = json.loads((tmp_path / "gouvernance" / "decision.json").read_text(encoding="utf-8"))
    assert saved["cadre"] == "TOGAF" and saved["recommandation_run_id"] == "RUN-1"


def test_decision_refusee_sans_notion_ou_sans_auteur(tmp_path):
    try:
        GouvernanceService(None, "Neyl", output_dir=tmp_path).enregistrer("ITIL", None, [], None)
        raise AssertionError("aurait dû refuser")
    except DecisionRefusee as e:
        assert "Notion" in str(e)
    try:
        GouvernanceService(FakeNotion(), "", output_dir=tmp_path).enregistrer("ITIL", None, [], None)
        raise AssertionError("aurait dû refuser")
    except DecisionRefusee as e:
        assert "Validateur" in str(e)


def test_matrices_pandas_et_classeur_xlsx():
    from openpyxl import load_workbook
    from io import BytesIO

    exigence = Exigence.model_validate({**EXEMPLE_SPEC, "processus_impactes": ["Gestion des risques"], "couverture": "Partielle", "couverture_detail": [{"controle_id": "CT-0001", "couverture": "partielle", "ecart": "PSSI incomplète"}]})
    texte = Texte(id="TX-0001", source="eurlex", reference="CELEX:32022R2554", titre="DORA", type="Règlement UE", url_officielle="https://eur-lex.europa.eu/x")
    tables = toutes([exigence], {"TX-0001": texte}, {}, [], [])
    assert list(tables) == ["M1", "M2", "M3", "M4", "M5"]
    assert tables["M1"].iloc[0]["ID exigence"] == "EX-0007" and tables["M1"].iloc[0]["Texte"] == "TX-0001 · DORA"
    assert tables["M2"].iloc[0]["Processus"] == "Gestion des risques"
    assert tables["M3"].iloc[0]["Écart"] == "CT-0001 : PSSI incomplète"
    assert tables["M4"].empty and tables["M5"].empty
    wb = load_workbook(BytesIO(classeur_xlsx(tables)))
    assert wb.sheetnames == ["M1 Traçabilité", "M2 Impact", "M3 Couverture", "M4 Risque", "M5 Jurisprudence"]
    assert wb["M1 Traçabilité"]["A1"].value == "ID exigence" and wb["M1 Traçabilité"].freeze_panes == "A2"
